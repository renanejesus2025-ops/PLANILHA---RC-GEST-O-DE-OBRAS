"""
Estilos visuais mínimos para o workbook gerado na Etapa 2, estendidos na
Etapa 4 (formato de moeda brasileira, ver `FORMATO_MOEDA_BR`).

Usa SOMENTE as 4 cores oficiais da marca, já transcritas em
`03_DESIGN_E_IDENTIDADE_VISUAL/DSG_001_...md` Seção A.5 (nenhuma cor foi
inventada; nenhuma cor semântica de alerta foi usada, porque o próprio
DSG_001 registra que elas ainda não foram homologadas — Seção C.3).

Tipografia: a fonte oficial "Gravesend Sans" (DSG_001 A.6) NÃO é aplicada
aqui porque o fallback para Excel continua [H] (DSG_001 Seção C.2) — usar
a fonte-padrão do Excel evita quebra silenciosa da identidade visual em
máquinas sem a fonte instalada. Quando o fallback for homologado, esta
é a única peça a ajustar (Font.name abaixo).
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Cores oficiais (DSG_001 Seção A.5) — formato ARGB usado pelo openpyxl.
BRANCO = "FFFFFFFF"
CINZA_AZULADO_CLARO = "FFABB6BA"
CINZA_ESVERDEADO_ESCURO = "FF5A5F5B"
AZUL_MARINHO_ESCURO = "FF0F1C2C"

FONTE_CABECALHO = Font(bold=True, color=BRANCO)
PREENCHIMENTO_CABECALHO = PatternFill("solid", fgColor=AZUL_MARINHO_ESCURO)
ALINHAMENTO_CABECALHO = Alignment(horizontal="center", vertical="center", wrap_text=True)

_BORDA_FINA = Side(style="thin", color=CINZA_AZULADO_CLARO)
BORDA_CELULA = Border(left=_BORDA_FINA, right=_BORDA_FINA, top=_BORDA_FINA, bottom=_BORDA_FINA)

FONTE_ROTULO_FORM = Font(bold=True, color=CINZA_ESVERDEADO_ESCURO)

# Diferenciação visual entre célula de entrada e célula calculada
# (AGENTS.md §10: "células de entrada claramente identificadas"; Etapa 3
# Seção 24). Usa a mesma cor neutra clara da marca, sem inventar uma cor
# nova — apenas um preenchimento leve para sinalizar "não digite aqui,
# isto é calculado por fórmula".
PREENCHIMENTO_CALCULADO = PatternFill("solid", fgColor=CINZA_AZULADO_CLARO)
FONTE_CALCULADA = Font(italic=True, color=CINZA_ESVERDEADO_ESCURO)

# Formatação monetária brasileira ("R$ 0,00") — Etapa 4, Seção 27,
# exigência explícita e literal desta homologação. Usada apenas nas
# abas novas desta etapa (Financeiro/Pagamentos/Resumo Financeiro); as
# colunas monetárias das Etapas 2/3 (Etapas/Subetapas/Serviços/Início)
# continuam com o formato "#,##0.00" já testado/homologado nessas
# etapas — não foram alteradas retroativamente (fora do escopo desta
# etapa, ver relatório da Etapa 4, Limitações). Os separadores de milhar/
# decimal exibidos seguem a configuração regional do Excel do usuário —
# o literal "R$" é a única parte fixa deste formato.
FORMATO_MOEDA_BR = '"R$" #,##0.00'
