"""
Testes do construtor do primeiro arquivo Excel funcional (Etapa 2).

Cobre os 10 itens mínimos da Seção 18 da homologação da Etapa 2:
1. arquivo Excel é criado;
2. arquivo pode ser aberto pelo openpyxl;
3. abas esperadas existem;
4. estrutura de cabeçalhos existe;
5. IDs são únicos;
6. hierarquia Obra → Etapa → Subetapa → Serviço é válida;
7. serviço não pode existir sem Subetapa;
8. Subetapa não pode existir sem Etapa;
9. dropdowns/listas amigáveis não expõem IDs técnicos como rótulos;
10. arquivo não apresenta corrupção estrutural.

Como não há, neste ambiente, um motor de planilha (Excel/LibreOffice)
disponível para avaliar fórmulas (ver relatório da Etapa 2), a
verificação de IDs/hierarquia é feita em dois níveis complementares:
(a) nas linhas literais, escritas com valores já resolvidos pelo
`GeradorId`/`BaseDados` da Etapa 1; (b) nas linhas-modelo em branco,
verificando que a fórmula de cada linha é, por construção, distinta e
sequencial (o padrão embute `ROW()`, garantindo unicidade quando calculado).
"""

from __future__ import annotations

import re
import zipfile
import xml.etree.ElementTree as ET
from datetime import date

import openpyxl
import pytest

from src.base_dados.repositorio import BaseDados
from src.excecoes import ErroReferenciaInvalida
from src.excel.construtor_workbook import (
    COLUNAS_ETAPAS,
    COLUNAS_SERVICOS,
    COLUNAS_SUBETAPAS,
    LINHAS_MODELO,
    construir_workbook,
)
from src.ids.gerador_id import eh_formato_de_id
from src.modelo.entidades import Etapa, Obra, ServicoOrcamento, Subetapa

ABAS_ESPERADAS = [
    "Início",
    "Etapas",
    "Subetapas",
    "Serviços",
    "Financeiro",
    "Pagamentos",
    "Resumo Financeiro",
    "Base_Dados",
]
# ^ Atualizado na Etapa 4 (Financeiro) — 3 abas novas ("Financeiro",
# "Pagamentos", "Resumo Financeiro") acrescentadas entre "Serviços" e
# "Base_Dados"; nenhuma aba da Etapa 2/3 foi removida ou renomeada.
# Mesmo tipo de ajuste já registrado na Etapa 3 para `test_entidades.py`/
# `test_excel_construtor.py` (relatório da Etapa 3, Seção 3).


def _construir_base_de_exemplo() -> BaseDados:
    """Uma obra pequena, com 2 Etapas, 3 Subetapas e 3 Serviços."""
    base = BaseDados()
    gid = base.gerador_id

    obra = base.adicionar_obra(
        Obra(id=gid.gerar("OBRAS"), nome="Residencial Jardim das Flores")
    )
    etapa_alvenaria = base.adicionar_etapa(
        Etapa(id=gid.gerar("ETAPAS"), id_obra=obra.id, nome="Alvenaria", ordem=1)
    )
    etapa_eletrica = base.adicionar_etapa(
        Etapa(id=gid.gerar("ETAPAS"), id_obra=obra.id, nome="Elétrica", ordem=2)
    )
    sub_chapisco = base.adicionar_subetapa(
        Subetapa(id=gid.gerar("SUBETAPAS"), id_etapa=etapa_alvenaria.id, nome="Chapisco")
    )
    sub_reboco = base.adicionar_subetapa(
        Subetapa(id=gid.gerar("SUBETAPAS"), id_etapa=etapa_alvenaria.id, nome="Reboco")
    )
    sub_fiacao = base.adicionar_subetapa(
        Subetapa(id=gid.gerar("SUBETAPAS"), id_etapa=etapa_eletrica.id, nome="Fiação")
    )
    base.adicionar_servico(
        ServicoOrcamento(
            id=gid.gerar("SERVICOS_ORCAMENTO"),
            id_subetapa=sub_chapisco.id,
            descricao="Chapisco de fachada",
            unidade_medida="m²",
            quantidade_orcada=200.0,
            valor_unitario=12.0,
        )
    )
    base.adicionar_servico(
        ServicoOrcamento(
            id=gid.gerar("SERVICOS_ORCAMENTO"),
            id_subetapa=sub_reboco.id,
            descricao="Reboco interno",
            unidade_medida="m²",
            quantidade_orcada=350.0,
            valor_unitario=18.5,
        )
    )
    base.adicionar_servico(
        ServicoOrcamento(
            id=gid.gerar("SERVICOS_ORCAMENTO"),
            id_subetapa=sub_fiacao.id,
            descricao="Fiação elétrica embutida",
            unidade_medida="m",
            quantidade_orcada=500.0,
            valor_unitario=7.2,
        )
    )
    return base


# item 1 e 2: arquivo é criado e pode ser reaberto pelo openpyxl -----------
def test_arquivo_e_criado_e_pode_ser_reaberto_pelo_openpyxl(tmp_path):
    caminho = tmp_path / "modelo.xlsx"
    wb = construir_workbook(BaseDados())
    wb.save(caminho)

    assert caminho.exists()
    assert caminho.stat().st_size > 0

    reaberto = openpyxl.load_workbook(caminho)
    assert reaberto is not None


# item 3: abas esperadas existem -------------------------------------------
def test_abas_esperadas_existem(tmp_path):
    caminho = tmp_path / "modelo.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    assert wb.sheetnames == ABAS_ESPERADAS


def test_aba_base_dados_fica_oculta(tmp_path):
    caminho = tmp_path / "modelo.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    assert wb["Base_Dados"].sheet_state == "hidden"


# item 4: estrutura de cabeçalhos existe ------------------------------------
@pytest.mark.parametrize(
    "aba,colunas",
    [("Etapas", COLUNAS_ETAPAS), ("Subetapas", COLUNAS_SUBETAPAS), ("Serviços", COLUNAS_SERVICOS)],
)
def test_cabecalhos_das_tabelas_existem(tmp_path, aba, colunas):
    caminho = tmp_path / "modelo.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    ws = wb[aba]
    titulos_esperados = [titulo for titulo, _largura, _oculta, _calculada in colunas]
    titulos_no_arquivo = [ws.cell(row=1, column=i + 1).value for i in range(len(colunas))]
    assert titulos_no_arquivo == titulos_esperados


def test_colunas_tecnicas_de_id_ficam_ocultas(tmp_path):
    caminho = tmp_path / "modelo.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    for aba, colunas in [
        ("Etapas", COLUNAS_ETAPAS),
        ("Subetapas", COLUNAS_SUBETAPAS),
        ("Serviços", COLUNAS_SERVICOS),
    ]:
        ws = wb[aba]
        for indice, (_titulo, _largura, oculta, _calculada) in enumerate(colunas, start=1):
            letra = ws.cell(row=1, column=indice).column_letter
            assert ws.column_dimensions[letra].hidden is oculta


# item 5: IDs são únicos ------------------------------------------------------
def test_ids_literais_sao_unicos_no_arquivo_gerado(tmp_path):
    base = _construir_base_de_exemplo()
    caminho = tmp_path / "modelo.xlsx"
    construir_workbook(base).save(caminho)
    wb = openpyxl.load_workbook(caminho)

    todos_os_ids: list[str] = []
    for aba, coluna_id in [("Etapas", 1), ("Subetapas", 1), ("Serviços", 1)]:
        ws = wb[aba]
        n_linhas_literais = {
            "Etapas": len(base.etapas),
            "Subetapas": len(base.subetapas),
            "Serviços": len(base.servicos),
        }[aba]
        for linha in range(2, 2 + n_linhas_literais):
            valor = ws.cell(row=linha, column=coluna_id).value
            assert isinstance(valor, str) and eh_formato_de_id(valor)
            todos_os_ids.append(valor)

    assert len(todos_os_ids) == len(set(todos_os_ids))


def test_formulas_de_id_das_linhas_modelo_sao_unicas_por_construcao(tmp_path):
    """
    As linhas em branco do modelo usam uma fórmula com ROW() distinto por
    linha — verificamos que a string da fórmula (logo, o ID que ela vai
    calcular) é diferente em cada linha, garantindo unicidade quando o
    Excel/Google Sheets efetivamente calcular a planilha.
    """
    caminho = tmp_path / "modelo.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    ws = wb["Etapas"]

    formulas = [ws.cell(row=r, column=1).value for r in range(2, 2 + LINHAS_MODELO)]
    assert len(formulas) == len(set(formulas))  # nenhuma fórmula duplicada
    for r, formula in zip(range(2, 2 + LINHAS_MODELO), formulas):
        assert f"ROW()-1" in formula
        assert f"$B{r}" in formula  # cada linha depende da própria linha, não de outra


# item 6: hierarquia Obra -> Etapa -> Subetapa -> Serviço é válida -----------
def test_hierarquia_completa_e_refletida_no_arquivo_com_nomes_amigaveis(tmp_path):
    """
    A partir da Etapa 3, as colunas de vínculo técnico (ID_Etapa/
    ID_Subetapa) são SEMPRE fórmula — inclusive nas linhas com dado
    literal (Seção "Uso do GeradorId" do módulo `construtor_workbook`) —
    para que continuem corretas se o Operador editar o dropdown de
    Etapa/Subetapa depois, diretamente no Excel. Por isso a verificação
    aqui é sobre o TEXTO da fórmula (aponta para a linha certa e para os
    intervalos certos), não sobre um valor literal — não há motor de
    planilha neste ambiente para calculá-la (ver relatório da Etapa 2).
    """
    base = _construir_base_de_exemplo()
    caminho = tmp_path / "modelo.xlsx"
    construir_workbook(base).save(caminho)
    wb = openpyxl.load_workbook(caminho)

    ws_sub = wb["Subetapas"]
    ws_srv = wb["Serviços"]

    # Subetapa "Reboco" (linha 3 dos dados de exemplo) deve mostrar a
    # Etapa amiga "Alvenaria" na coluna visível (C), e uma fórmula de
    # vínculo técnico coerente na coluna oculta (E).
    linha_reboco = 1 + list(base.subetapas.values()).index(
        next(s for s in base.subetapas.values() if s.nome == "Reboco")
    ) + 1
    assert ws_sub.cell(row=linha_reboco, column=3).value == "Alvenaria"
    formula_id_etapa = ws_sub.cell(row=linha_reboco, column=5).value
    assert f"$C{linha_reboco}" in formula_id_etapa
    assert "Etapas!$A:$A" in formula_id_etapa and "Etapas!$B:$B" in formula_id_etapa

    # Serviço "Reboco interno" deve mostrar a Subetapa amiga "Reboco" na
    # coluna visível (C), e uma fórmula de vínculo técnico coerente na
    # coluna oculta ID_Subetapa (M — Etapa 3 acrescentou colunas de
    # orçamento entre C e a antiga posição do vínculo técnico).
    linha_servico = 1 + list(base.servicos.values()).index(
        next(s for s in base.servicos.values() if s.descricao == "Reboco interno")
    ) + 1
    assert ws_srv.cell(row=linha_servico, column=3).value == "Reboco"
    formula_id_subetapa = ws_srv.cell(row=linha_servico, column=13).value
    assert f"$C{linha_servico}" in formula_id_subetapa
    assert "Subetapas!$A:$A" in formula_id_subetapa and "Subetapas!$B:$B" in formula_id_subetapa


# item 7 e 8: Serviço não pode existir sem Subetapa; Subetapa sem Etapa -----
def test_servico_nao_pode_ser_adicionado_a_base_sem_subetapa_valida():
    """
    Garante que a base que alimenta o gerador de Excel nunca aceita um
    Serviço órfão — logo, o arquivo gerado nunca poderá conter uma linha
    de Serviço sem Subetapa válida.
    """
    base = BaseDados()
    with pytest.raises(ErroReferenciaInvalida):
        base.adicionar_servico(
            ServicoOrcamento(id="SRV-0001", id_subetapa="SUB-9999", descricao="Reboco")
        )
    assert base.servicos == {}


def test_subetapa_nao_pode_ser_adicionada_a_base_sem_etapa_valida():
    base = BaseDados()
    with pytest.raises(ErroReferenciaInvalida):
        base.adicionar_subetapa(Subetapa(id="SUB-0001", id_etapa="ETA-9999", nome="Reboco"))
    assert base.subetapas == {}


# item 9: dropdowns não expõem IDs técnicos como rótulo ----------------------
def test_dropdowns_apontam_para_coluna_de_nome_nao_para_coluna_de_id(tmp_path):
    caminho = tmp_path / "modelo.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)

    lista_etapas = wb.defined_names["Lista_Etapas"].attr_text
    lista_subetapas = wb.defined_names["Lista_Subetapas"].attr_text
    # Coluna B é "Nome" nas duas abas (coluna A, técnica, fica oculta).
    assert re.search(r"Etapas!\$B\$", lista_etapas)
    assert re.search(r"Subetapas!\$B\$", lista_subetapas)
    assert "$A$" not in lista_etapas
    assert "$A$" not in lista_subetapas


def test_valores_exibidos_nas_colunas_de_dropdown_nao_sao_ids_tecnicos(tmp_path):
    base = _construir_base_de_exemplo()
    caminho = tmp_path / "modelo.xlsx"
    construir_workbook(base).save(caminho)
    wb = openpyxl.load_workbook(caminho)

    ws_sub = wb["Subetapas"]
    for linha in range(2, 2 + len(base.subetapas)):
        valor_coluna_etapa = ws_sub.cell(row=linha, column=3).value
        assert not eh_formato_de_id(valor_coluna_etapa)

    ws_srv = wb["Serviços"]
    for linha in range(2, 2 + len(base.servicos)):
        valor_coluna_subetapa = ws_srv.cell(row=linha, column=3).value
        assert not eh_formato_de_id(valor_coluna_subetapa)


# item 10: arquivo não apresenta corrupção estrutural ------------------------
def test_arquivo_gerado_nao_apresenta_corrupcao_estrutural(tmp_path):
    base = _construir_base_de_exemplo()
    caminho = tmp_path / "modelo.xlsx"
    construir_workbook(base).save(caminho)

    with zipfile.ZipFile(caminho) as z:
        assert z.testzip() is None  # nenhum membro do zip corrompido
        partes_xml = [n for n in z.namelist() if n.endswith(".xml") or n.endswith(".rels")]
        assert len(partes_xml) > 0
        for nome in partes_xml:
            ET.fromstring(z.read(nome))  # levanta ParseError se malformado


def test_reabertura_com_data_only_nao_gera_erro(tmp_path):
    """Reabrir com data_only=True (modo somente-valores) não deve falhar,
    mesmo sem um motor de planilha para calcular as fórmulas."""
    caminho = tmp_path / "modelo.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho, data_only=True)
    assert wb.sheetnames == ABAS_ESPERADAS
