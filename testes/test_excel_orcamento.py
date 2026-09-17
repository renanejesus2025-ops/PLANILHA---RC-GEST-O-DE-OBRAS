"""
Testes das fórmulas de orçamento adicionadas na aba Serviços/Subetapas/
Etapas/Início pela Etapa 3 ("Orçamento").

Como não há motor de planilha neste ambiente (ver relatório da Etapa 2 e
Seção 28 da Etapa 3), a verificação é sobre o TEXTO das fórmulas (padrão
correto, referências corretas) — não sobre um valor calculado. Os itens
18–20 da Seção 27 (workbook criado, reaberto pelo openpyxl, XML íntegro)
já estão cobertos por `test_excel_construtor.py`; aqui focamos no que é
específico da Etapa 3.
"""

from __future__ import annotations

import openpyxl
import pytest

from src.base_dados.repositorio import BaseDados
from src.excel.construtor_workbook import (
    COL_SRV_AJUSTE_MANUAL,
    COL_SRV_ID_SUBETAPA,
    COL_SRV_STATUS,
    COL_SRV_VALOR_CALCULADO,
    COL_SRV_VALOR_PREVISTO,
    COL_SRV_VALOR_PREVISTO_MANUAL,
    COL_SRV_VARIACAO,
    COLUNAS_SERVICOS,
    LINHAS_MODELO,
    construir_workbook,
)
from src.ids.gerador_id import eh_formato_de_id
from src.modelo.entidades import Etapa, Obra, ServicoOrcamento, Subetapa
from src.modelo.enums import StatusServico


def _base_com_um_servico(**kwargs) -> BaseDados:
    base = BaseDados()
    gid = base.gerador_id
    obra = base.adicionar_obra(Obra(id=gid.gerar("OBRAS"), nome="Obra"))
    etapa = base.adicionar_etapa(Etapa(id=gid.gerar("ETAPAS"), id_obra=obra.id, nome="Alvenaria"))
    subetapa = base.adicionar_subetapa(Subetapa(id=gid.gerar("SUBETAPAS"), id_etapa=etapa.id, nome="Reboco"))
    base.adicionar_servico(
        ServicoOrcamento(
            id=gid.gerar("SERVICOS_ORCAMENTO"),
            id_subetapa=subetapa.id,
            descricao="Reboco interno",
            **kwargs,
        )
    )
    return base


def test_valor_calculado_e_uma_formula_de_multiplicacao(tmp_path):
    base = _base_com_um_servico(quantidade_orcada=100, valor_unitario=50)
    caminho = tmp_path / "v.xlsx"
    construir_workbook(base).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Serviços"]
    formula = ws.cell(row=2, column=COL_SRV_VALOR_CALCULADO).value
    assert "$E2*$F2" in formula  # Quantidade (E) × Valor Unitário (F)


def test_ajuste_manual_e_uma_formula_se_entao(tmp_path):
    base = _base_com_um_servico(quantidade_orcada=100, valor_unitario=50)
    caminho = tmp_path / "v.xlsx"
    construir_workbook(base).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Serviços"]
    formula = ws.cell(row=2, column=COL_SRV_AJUSTE_MANUAL).value
    assert '"NÃO"' in formula and '"SIM"' in formula
    assert "$H2" in formula  # depende da coluna Valor Previsto (Ajuste Manual)


def test_valor_previsto_prioriza_o_ajuste_manual_quando_preenchido(tmp_path):
    base = _base_com_um_servico(quantidade_orcada=100, valor_unitario=50)
    caminho = tmp_path / "v.xlsx"
    construir_workbook(base).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Serviços"]
    formula = ws.cell(row=2, column=COL_SRV_VALOR_PREVISTO).value
    # =IF($H2="",$G2,$H2)  -> usa Valor Calculado (G) se H estiver vazio, senão usa H
    assert formula == '=IF($H2="",$G2,$H2)'


def test_variacao_e_previsto_menos_calculado(tmp_path):
    base = _base_com_um_servico(quantidade_orcada=100, valor_unitario=50)
    caminho = tmp_path / "v.xlsx"
    construir_workbook(base).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Serviços"]
    formula = ws.cell(row=2, column=COL_SRV_VARIACAO).value
    assert "$J2-$G2" in formula  # Valor Previsto (J) - Valor Calculado (G)
    assert "Desvio" not in formula and "desvio" not in formula  # termo proibido (Seção 13)


def test_valor_previsto_manual_e_escrito_como_valor_literal_quando_ha_ajuste(tmp_path):
    base = _base_com_um_servico(quantidade_orcada=100, valor_unitario=50, valor_previsto_manual=5300)
    caminho = tmp_path / "v.xlsx"
    construir_workbook(base).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Serviços"]
    assert ws.cell(row=2, column=COL_SRV_VALOR_PREVISTO_MANUAL).value == 5300


def test_dropdown_de_status_usa_exatamente_os_5_rotulos_homologados(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Serviços"]
    dvs_status = [dv for dv in ws.data_validations.dataValidation if str(dv.sqref).startswith("L")]
    assert len(dvs_status) == 1
    formula = dvs_status[0].formula1
    for status in StatusServico:
        assert status.rotulo in formula
    # nenhum rótulo é, por si, um ID técnico
    for status in StatusServico:
        assert not eh_formato_de_id(status.rotulo)


def test_status_literal_escrito_com_rotulo_amigavel_nao_com_valor_do_enum(tmp_path):
    base = _base_com_um_servico(status=StatusServico.CONCLUIDO)
    caminho = tmp_path / "v.xlsx"
    construir_workbook(base).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Serviços"]
    assert ws.cell(row=2, column=COL_SRV_STATUS).value == "Concluído"


def test_total_previsto_da_subetapa_referencia_valor_previsto_e_exclui_status(tmp_path):
    base = _base_com_um_servico(quantidade_orcada=100, valor_unitario=50)
    caminho = tmp_path / "v.xlsx"
    construir_workbook(base).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    ws_sub = wb["Subetapas"]
    formula = ws_sub.cell(row=2, column=6).value  # coluna F: Total Previsto
    assert "SUMIFS(" in formula
    assert "Serviços!$J:$J" in formula  # soma o Valor Previsto (J), não o Valor Calculado
    assert "Serviços!$M:$M" in formula  # casa pelo ID_Subetapa técnico (M) — nunca pelo nome
    assert "<>Cancelado" in formula
    assert "<>Retirado do Escopo" in formula
    assert "<>Substituído" in formula


def test_total_previsto_da_etapa_soma_totais_das_subetapas(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws_eta = openpyxl.load_workbook(caminho)["Etapas"]
    formula = ws_eta.cell(row=2, column=6).value  # coluna F: Total Previsto
    assert "SUMIFS(Subetapas!$F:$F" in formula
    assert "Subetapas!$E:$E" in formula  # casa pelo ID_Etapa técnico (E) da Subetapa


def test_orcamento_inicial_soma_os_totais_das_etapas(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws_ini = openpyxl.load_workbook(caminho)["Início"]
    formula = ws_ini.cell(row=8, column=2).value
    assert formula.startswith("=SUM(Etapas!$F$2:$F$")


def test_totais_nao_somam_aportes_alteracoes_ou_custo_realizado(tmp_path):
    """
    Seção 14/15/16/17: nenhuma dessas camadas deve aparecer nas fórmulas
    de total desta etapa.
    """
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    for aba in ("Início", "Etapas", "Subetapas", "Serviços"):
        ws = wb[aba]
        for linha in ws.iter_rows():
            for celula in linha:
                if isinstance(celula.value, str) and celula.value.startswith("="):
                    for termo_proibido in ("Aporte", "Alteracao", "Alteração", "Custo Realizado", "Saldo"):
                        assert termo_proibido not in celula.value


def test_quantidade_e_valor_unitario_tem_validacao_numerica_maior_ou_igual_a_zero(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Serviços"]
    validacoes = {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}
    dv_quantidade = next(dv for ref, dv in validacoes.items() if ref.startswith("E2"))
    dv_valor = next(dv for ref, dv in validacoes.items() if ref.startswith("F2"))
    assert dv_quantidade.type == "decimal" and dv_quantidade.operator == "greaterThanOrEqual"
    assert dv_quantidade.formula1 == "0"
    assert dv_valor.type == "decimal" and dv_valor.operator == "greaterThanOrEqual"
    assert dv_valor.formula1 == "0"


def test_colunas_calculadas_tem_estilo_visualmente_distinto_de_entrada(tmp_path):
    """
    AGENTS.md §10: "células de entrada claramente identificadas". As
    colunas calculadas (G, I, J, K) devem ter um preenchimento diferente
    da célula de entrada (E, F, H).
    """
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Serviços"]
    preenchimento_calculado = ws.cell(row=2, column=COL_SRV_VALOR_CALCULADO).fill.fgColor.rgb
    preenchimento_entrada = ws.cell(row=2, column=5).fill.fgColor.rgb  # Quantidade Orçada (entrada)
    assert preenchimento_calculado != preenchimento_entrada


def test_autofilter_existe_na_aba_servicos(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Serviços"]
    assert ws.auto_filter.ref is not None


def test_nenhuma_formatacao_condicional_de_alerta_foi_criada(tmp_path):
    """Seção 31: não implementar alertas ainda (ex.: cor na Variação)."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Serviços"]
    assert len(ws.conditional_formatting._cf_rules) == 0
