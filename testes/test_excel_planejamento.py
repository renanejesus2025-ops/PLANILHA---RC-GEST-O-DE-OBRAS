"""
Testes das colunas de Cronograma acrescentadas às abas Etapas e
Subetapas, e do cronograma consolidado da Obra na aba Início (Etapa 8).

Mesmo raciocínio das etapas anteriores: sem motor de planilha neste
ambiente — verificação pelo TEXTO da fórmula e pela estrutura.

Registro da decisão de layout: o cronograma NÃO ganhou uma aba própria.
As datas moram nas próprias Etapas/Subetapas (decisão homologada em
2026-09-23, Opção A) e, por isso, ficam nas abas dessas entidades — que
já são Tabelas Excel estruturadas e se expandem sozinhas. Uma terceira
aba "Cronograma" montada por fórmulas teria buffer fixo de linhas e
voltaria a ignorar registros além dele (defeito AUD-20/AUD-21, corrigido
na Etapa 5.1).
"""

from __future__ import annotations

from datetime import date

import openpyxl
from openpyxl.utils import get_column_letter

from src.base_dados.repositorio import BaseDados
from src.excel.construtor_workbook import (
    COL_ETA_DATA_FIM_PREVISTA,
    COL_ETA_DATA_FIM_REAL,
    COL_ETA_DATA_INICIO_PREVISTA,
    COL_ETA_DATA_INICIO_REAL,
    COL_ETA_DURACAO_PREVISTA,
    COL_ETA_DURACAO_REAL,
    COL_ETA_VARIACAO_PRAZO,
    COL_SUB_DATA_INICIO_PREVISTA,
    COL_SUB_DURACAO_PREVISTA,
    COL_SUB_VARIACAO_PRAZO,
    COLUNAS_ETAPAS,
    COLUNAS_SUBETAPAS,
    LINHAS_MODELO,
    construir_workbook,
)
from src.modelo.entidades import Etapa, Obra, Subetapa

TITULOS_CRONOGRAMA = [
    "Data Início Prevista",
    "Data Fim Prevista",
    "Data Início Real",
    "Data Fim Real",
    "Duração Prevista (dias)",
    "Duração Real (dias)",
    "Variação de Prazo (dias)",
]


def _base_com_cronograma() -> BaseDados:
    base = BaseDados()
    obra = base.adicionar_obra(Obra(id=base.gerador_id.gerar("OBRAS"), nome="Obra"))
    etapa = base.adicionar_etapa(
        Etapa(
            id=base.gerador_id.gerar("ETAPAS"), id_obra=obra.id, nome="Alvenaria",
            data_inicio_prevista=date(2026, 3, 1), data_fim_prevista=date(2026, 4, 30),
            data_inicio_real=date(2026, 3, 5), data_fim_real=date(2026, 5, 10),
        )
    )
    base.adicionar_subetapa(
        Subetapa(
            id=base.gerador_id.gerar("SUBETAPAS"), id_etapa=etapa.id, nome="Reboco",
            data_inicio_prevista=date(2026, 3, 1), data_fim_prevista=date(2026, 3, 15),
        )
    )
    return base


# ----------------------------------------------------------------------
# Estrutura das colunas novas
# ----------------------------------------------------------------------
def test_etapas_ganhou_as_7_colunas_de_cronograma(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Etapas"]
    titulos = [ws.cell(row=1, column=i + 1).value for i in range(len(COLUNAS_ETAPAS))]
    assert titulos == [t for t, _l, _o, _c in COLUNAS_ETAPAS]
    assert titulos[-7:] == TITULOS_CRONOGRAMA


def test_subetapas_ganhou_as_mesmas_7_colunas_de_cronograma(tmp_path):
    """SUBETAPAS não tinha nenhum campo de data antes da Etapa 8."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Subetapas"]
    titulos = [ws.cell(row=1, column=i + 1).value for i in range(len(COLUNAS_SUBETAPAS))]
    assert titulos == [t for t, _l, _o, _c in COLUNAS_SUBETAPAS]
    assert titulos[-7:] == TITULOS_CRONOGRAMA


def test_nenhuma_aba_nova_foi_criada_na_etapa_8(tmp_path):
    """O cronograma mora nas abas das próprias entidades — nenhuma aba
    "Cronograma"/"Planejamento" com buffer fixo foi criada (AUD-20/21)."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    assert "Planejamento" not in wb.sheetnames
    assert "Cronograma" not in wb.sheetnames
    assert "Cronograma da Obra" not in wb.sheetnames


def test_tabelas_estruturadas_cobrem_as_colunas_novas(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    ultima_eta = get_column_letter(len(COLUNAS_ETAPAS))
    ultima_sub = get_column_letter(len(COLUNAS_SUBETAPAS))
    assert wb["Etapas"].tables["TabelaEtapas"].ref == f"A1:{ultima_eta}{1 + LINHAS_MODELO}"
    assert wb["Subetapas"].tables["TabelaSubetapas"].ref == f"A1:{ultima_sub}{1 + LINHAS_MODELO}"


# ----------------------------------------------------------------------
# Fórmulas calculadas
# ----------------------------------------------------------------------
def test_formula_duracao_prevista_subtrai_datas(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Etapas"]
    letra_ini = get_column_letter(COL_ETA_DATA_INICIO_PREVISTA)
    letra_fim = get_column_letter(COL_ETA_DATA_FIM_PREVISTA)
    formula = ws.cell(row=2, column=COL_ETA_DURACAO_PREVISTA).value
    assert formula == (
        f'=IF(OR(${letra_ini}2="",${letra_fim}2=""),"",${letra_fim}2-${letra_ini}2)'
    )
    assert "+1" not in formula  # subtração pura, sem contagem inclusiva


def test_formula_duracao_real_subtrai_datas_reais(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Etapas"]
    letra_ini = get_column_letter(COL_ETA_DATA_INICIO_REAL)
    letra_fim = get_column_letter(COL_ETA_DATA_FIM_REAL)
    formula = ws.cell(row=2, column=COL_ETA_DURACAO_REAL).value
    assert f"${letra_fim}2-${letra_ini}2" in formula


def test_formula_variacao_de_prazo_e_real_menos_previsto(tmp_path):
    """Positivo = atraso; negativo = antecipação (convenção de sinal já
    homologada em `Alteracao.impacto_prazo_dias`)."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Etapas"]
    letra_fim_prev = get_column_letter(COL_ETA_DATA_FIM_PREVISTA)
    letra_fim_real = get_column_letter(COL_ETA_DATA_FIM_REAL)
    formula = ws.cell(row=2, column=COL_ETA_VARIACAO_PRAZO).value
    assert f"${letra_fim_real}2-${letra_fim_prev}2" in formula


def test_subetapas_tem_as_mesmas_formulas_de_cronograma(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Subetapas"]
    assert "IF(OR(" in ws.cell(row=2, column=COL_SUB_DURACAO_PREVISTA).value
    assert "IF(OR(" in ws.cell(row=2, column=COL_SUB_VARIACAO_PRAZO).value


def test_colunas_calculadas_sao_formula_tambem_na_linha_com_dado_literal(tmp_path):
    """Precisam continuar corretas se o Operador editar as datas direto
    no Excel — mesmo princípio de todas as etapas anteriores."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(_base_com_cronograma()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Etapas"]
    assert str(ws.cell(row=2, column=COL_ETA_DURACAO_PREVISTA).value).startswith("=")
    assert str(ws.cell(row=2, column=COL_ETA_VARIACAO_PRAZO).value).startswith("=")


# ----------------------------------------------------------------------
# Datas literais e formatação
# ----------------------------------------------------------------------
def test_datas_literais_sao_escritas_na_linha_da_etapa(tmp_path):
    """O Excel não tem tipo "data sem hora": uma `date` gravada volta do
    arquivo como `datetime` à meia-noite. A comparação é feita sobre
    `.date()` — o dia gravado tem de ser exatamente o dia da entidade."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(_base_com_cronograma()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Etapas"]
    assert ws.cell(row=2, column=COL_ETA_DATA_INICIO_PREVISTA).value.date() == date(2026, 3, 1)
    assert ws.cell(row=2, column=COL_ETA_DATA_FIM_REAL).value.date() == date(2026, 5, 10)


def test_datas_da_subetapa_sao_escritas(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(_base_com_cronograma()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Subetapas"]
    assert ws.cell(row=2, column=COL_SUB_DATA_INICIO_PREVISTA).value.date() == date(2026, 3, 1)


def test_colunas_de_data_usam_formato_de_data_br(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Etapas"]
    for coluna in (
        COL_ETA_DATA_INICIO_PREVISTA, COL_ETA_DATA_FIM_PREVISTA,
        COL_ETA_DATA_INICIO_REAL, COL_ETA_DATA_FIM_REAL,
    ):
        assert ws.cell(row=2, column=coluna).number_format == "DD/MM/YYYY"


def test_colunas_de_dias_sao_numero_nunca_data(tmp_path):
    """Uma diferença de datas formatada como data exibiria "00/01/1900"
    em vez de "0 dias"."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Etapas"]
    for coluna in (COL_ETA_DURACAO_PREVISTA, COL_ETA_DURACAO_REAL, COL_ETA_VARIACAO_PRAZO):
        assert ws.cell(row=2, column=coluna).number_format == "0"


def test_colunas_de_data_tem_validacao_de_data(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Etapas"]
    for coluna in (
        COL_ETA_DATA_INICIO_PREVISTA, COL_ETA_DATA_FIM_PREVISTA,
        COL_ETA_DATA_INICIO_REAL, COL_ETA_DATA_FIM_REAL,
    ):
        letra = get_column_letter(coluna)
        dv = next(
            dv for dv in ws.data_validations.dataValidation
            if str(dv.sqref).startswith(f"{letra}2") and dv.type == "date"
        )
        assert dv.allow_blank
        # `"1900-01-01"` seria avaliado pelo Excel como a expressão
        # aritmética 1900−1−1 = 1898, e o limite viraria "serial ≥ 1898"
        # (≈ 13/03/1905) — defeito apontado pela auditoria da Etapa 8.
        assert dv.formula1 == "DATE(1900,1,1)"


# ----------------------------------------------------------------------
# Cronograma consolidado da Obra (aba Início)
# ----------------------------------------------------------------------
def _rotulos_do_inicio(caminho) -> dict:
    ws = openpyxl.load_workbook(caminho)["Início"]
    return {
        ws.cell(row=r, column=1).value: ws.cell(row=r, column=2)
        for r in range(1, 20)
        if ws.cell(row=r, column=1).value
    }


def test_inicio_tem_as_datas_consolidadas_da_obra(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    rotulos = _rotulos_do_inicio(caminho)
    assert "Início Previsto (menor data das Etapas)" in rotulos
    assert "Fim Previsto (maior data das Etapas)" in rotulos


def test_consolidacao_da_obra_usa_referencia_estruturada_sem_range_fixo(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    rotulos = _rotulos_do_inicio(caminho)
    formula_inicio = rotulos["Início Previsto (menor data das Etapas)"].value
    assert "MIN(TabelaEtapas[Data Início Prevista])" in formula_inicio
    assert "$" not in formula_inicio  # nenhuma referência de célula fixa
    formula_fim = rotulos["Fim Previsto (maior data das Etapas)"].value
    assert "MAX(TabelaEtapas[Data Fim Prevista])" in formula_fim


def test_excel_nao_tem_indicador_de_variacao_de_prazo_da_obra(tmp_path):
    """
    O indicador consolidado "Variação de Prazo da Obra" foi REJEITADO e
    REMOVIDO por decisão do responsável do projeto em 2026-09-23
    (Alternativa D1): não havia regra de negócio homologada que
    definisse seu significado, e a conta `MAX(Fim Real) − MAX(Fim
    Previsto)` podia tomar os dois operandos de Etapas diferentes.

    Este teste impede que ele volte por descuido: nenhuma célula da aba
    Início pode conter o rótulo nem uma subtração entre as duas colunas
    consolidadas. A Variação de Prazo permanece nas abas Etapas e
    Subetapas, por registro.
    """
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Início"]

    rotulos = [str(ws.cell(row=r, column=1).value or "") for r in range(1, 25)]
    assert not any("Variação de Prazo da Obra" in r for r in rotulos)

    for r in range(1, 25):
        valor = str(ws.cell(row=r, column=2).value or "")
        # Nenhuma subtração entre Fim Real e Fim Previsto consolidados.
        assert "TabelaEtapas[Data Fim Real])-MAX(" not in valor
        assert "ha_etapa_prevista" not in valor


def test_consolidacao_da_obra_trata_ausencia_de_datas_como_vazio(tmp_path):
    """Sem COUNT(), MIN/MAX retornariam 0 — exibido como 00/01/1900."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    rotulos = _rotulos_do_inicio(caminho)
    assert 'COUNT(TabelaEtapas[Data Início Prevista])=0,""' in rotulos[
        "Início Previsto (menor data das Etapas)"
    ].value


def test_datas_declaradas_da_obra_continuam_existindo_e_separadas(tmp_path):
    """A consolidação não substitui nem sobrescreve o que o Operador
    declarou para a Obra — os dois rótulos coexistem (REG-032)."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    rotulos = _rotulos_do_inicio(caminho)
    assert "Data Inicial Planejada" in rotulos  # declarada pelo Operador
    assert "Data Final Planejada" in rotulos
    assert "Início Previsto (menor data das Etapas)" in rotulos  # derivada


def test_orcamento_inicial_continua_na_celula_b8_referenciada_pelo_resumo(tmp_path):
    """Regressão: o Resumo Financeiro referencia `Início!$B$8`. As linhas
    novas da Etapa 8 entraram DEPOIS dessa célula, não antes."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Início"]
    assert ws.cell(row=8, column=1).value == "Orçamento Inicial (Previsto)"
