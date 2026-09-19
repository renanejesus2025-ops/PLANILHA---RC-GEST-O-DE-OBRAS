"""
Testes da aba Alterações e da atualização do Resumo Financeiro (Etapa 5).

Mesmo raciocínio de `testes/test_excel_financeiro.py` (Etapa 4): sem
motor de planilha neste ambiente por padrão — verificação pelo TEXTO da
fórmula. Ver também Task 4 deste plano para validação real via Excel COM
quando disponível.
"""

from __future__ import annotations

import openpyxl
import pytest

from src.base_dados.repositorio import BaseDados
from src.excel.construtor_workbook import (
    COL_ALT_IMPACTO_ORCAMENTO,
    COL_ALT_STATUS,
    COL_ALT_TIPO,
    COLUNAS_ALTERACOES,
    construir_workbook,
)
from src.modelo.enums import StatusAprovacaoAlteracao, TipoAlteracao


def test_aba_alteracoes_existe(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    assert "Alterações" in wb.sheetnames


def test_cabecalho_da_aba_alteracoes(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Alterações"]
    titulos_esperados = [titulo for titulo, _l, _o, _c in COLUNAS_ALTERACOES]
    titulos_no_arquivo = [ws.cell(row=1, column=i + 1).value for i in range(len(COLUNAS_ALTERACOES))]
    assert titulos_no_arquivo == titulos_esperados


def test_dropdown_de_tipo_de_alteracao_tem_3_valores(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Alterações"]
    letra_tipo = ws.cell(row=1, column=COL_ALT_TIPO).column_letter
    dvs = [dv for dv in ws.data_validations.dataValidation if str(dv.sqref).startswith(f"{letra_tipo}2")]
    assert len(dvs) == 1
    for tipo in TipoAlteracao:
        assert tipo.rotulo in dvs[0].formula1
    assert dvs[0].formula1.count(",") == 2  # exatamente 3 itens


def test_dropdown_de_status_de_aprovacao_tem_4_valores(tmp_path):
    """Domínio corrigido na Etapa 5.1: Em análise, Aprovada, Rejeitada, Cancelada."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Alterações"]
    letra_status = ws.cell(row=1, column=COL_ALT_STATUS).column_letter
    dvs = [dv for dv in ws.data_validations.dataValidation if str(dv.sqref).startswith(f"{letra_status}2")]
    assert len(dvs) == 1
    for status in StatusAprovacaoAlteracao:
        assert status.rotulo in dvs[0].formula1
    assert dvs[0].formula1.count(",") == 3  # exatamente 4 itens


def test_impacto_no_orcamento_aceita_negativo_sem_validacao_de_sinal(tmp_path):
    """Diferente de Financeiro!Valor (Etapa 4) — aqui NÃO há DataValidation de >= 0."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Alterações"]
    letra = ws.cell(row=1, column=COL_ALT_IMPACTO_ORCAMENTO).column_letter
    dvs_de_sinal = [
        dv for dv in ws.data_validations.dataValidation
        if str(dv.sqref).startswith(f"{letra}2") and dv.operator == "greaterThanOrEqual"
    ]
    assert dvs_de_sinal == []


def test_resumo_financeiro_orcamento_vigente_soma_3_parcelas(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Resumo Financeiro"]
    linhas = {ws.cell(row=r, column=1).value: r for r in range(1, 20)}
    linha_vigente = linhas["Orçamento Vigente"]
    formula = ws.cell(row=linha_vigente, column=2).value
    assert formula.count("+") == 2  # 3 parcelas = 2 somas
    assert "Alterações Aprovadas" in linhas


def test_resumo_financeiro_alteracoes_aprovadas_usa_sumifs_por_status(tmp_path):
    """Etapa 5.1 (AUD-21): referência estruturada de Tabela
    (`TabelaAlteracoes[...]`), sem prefixo de aba nem range fixo de
    linhas — estruturais em Excel não precisam de qualificador de aba."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Resumo Financeiro"]
    linhas = {ws.cell(row=r, column=1).value: r for r in range(1, 20)}
    linha = linhas["Alterações Aprovadas"]
    formula = ws.cell(row=linha, column=2).value
    assert "SUMIFS(" in formula
    assert "TabelaAlteracoes[" in formula
    assert StatusAprovacaoAlteracao.APROVADA.rotulo in formula
