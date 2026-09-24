"""
Testes da aba Compras (Etapa 6).

Mesmo raciocínio de `testes/test_excel_alteracoes.py` (Etapa 5): sem
motor de planilha neste ambiente — verificação pelo TEXTO da fórmula.
"""

from __future__ import annotations

import openpyxl

from src.base_dados.repositorio import BaseDados
from src.excel.construtor_workbook import (
    COL_COM_DESCRICAO,
    COL_COM_ETAPA,
    COL_COM_FORMA_PAGAMENTO,
    COL_COM_ID_SERVICO,
    COL_COM_PLANEJADO,
    COL_COM_SERVICO,
    COL_COM_STATUS_APROVACAO,
    COL_COM_STATUS_ENTREGA,
    COL_COM_SUBETAPA,
    COL_COM_VALOR,
    COL_COM_VARIACAO_PERCENTUAL,
    COL_COM_VARIACAO_REAIS,
    COLUNAS_COMPRAS,
    construir_workbook,
)


def test_aba_compras_existe(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    assert "Compras" in wb.sheetnames


def test_cabecalho_da_aba_compras(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Compras"]
    titulos_esperados = [titulo for titulo, _l, _o, _c in COLUNAS_COMPRAS]
    titulos_no_arquivo = [ws.cell(row=1, column=i + 1).value for i in range(len(COLUNAS_COMPRAS))]
    assert titulos_no_arquivo == titulos_esperados


def test_aba_compras_e_uma_tabela_estruturada_sem_range_fixo(tmp_path):
    """Etapa 6 integra ao padrão da Etapa 5.1 — expansão real, sem `LINHAS_MODELO` como teto."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Compras"]
    tabela = ws.tables["TabelaCompras"]
    assert tabela.ref == f"A1:S{1 + 30}"  # cabeçalho + LINHAS_MODELO (buffer inicial)
    assert tabela.autoFilter is not None


def test_dropdowns_de_etapa_subetapa_servico_sao_opcionais_e_independentes(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Compras"]

    letra_etapa = ws.cell(row=1, column=COL_COM_ETAPA).column_letter
    letra_subetapa = ws.cell(row=1, column=COL_COM_SUBETAPA).column_letter
    letra_servico = ws.cell(row=1, column=COL_COM_SERVICO).column_letter

    dv_etapa = next(dv for dv in ws.data_validations.dataValidation if str(dv.sqref).startswith(f"{letra_etapa}2"))
    dv_subetapa = next(dv for dv in ws.data_validations.dataValidation if str(dv.sqref).startswith(f"{letra_subetapa}2"))
    dv_servico = next(dv for dv in ws.data_validations.dataValidation if str(dv.sqref).startswith(f"{letra_servico}2"))

    assert dv_etapa.formula1 == "Lista_Etapas" and dv_etapa.allow_blank
    assert dv_subetapa.formula1 == "Lista_Subetapas" and dv_subetapa.allow_blank
    assert dv_servico.formula1 == "Lista_Servicos" and dv_servico.allow_blank


def test_lista_servicos_aponta_para_descricao_nao_para_id(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    lista = wb.defined_names["Lista_Servicos"].attr_text
    assert lista == "TabelaServicos[Descrição]"
    assert "ID" not in lista


def test_servico_tem_validacao_de_nome_unico_countif(tmp_path):
    """
    Etapa 6 estende AUD-18 (Etapa 5.1) à coluna Descrição de Serviços,
    agora que Compras também resolve o vínculo por nome.

    **Corrigido na validação Excel/COM da Etapa 8:** a asserção original
    exigia a referência estruturada `TabelaServicos[Descrição]` dentro da
    fórmula de validação — e era justamente isso que fazia o Excel real
    recusar o arquivo inteiro (nenhuma versão de V5 a V8 abria). A
    validação passa a usar o intervalo nomeado equivalente, que aponta
    para a MESMA coluna da MESMA tabela (sem range fixo).
    """
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Serviços"]
    dv_unico = next(dv for dv in ws.data_validations.dataValidation if str(dv.sqref).startswith("B2") and dv.type == "custom")
    assert "COUNTIF" in dv_unico.formula1
    assert "Lista_Servicos" in dv_unico.formula1


def test_nenhum_dropdown_inventado_para_forma_pagamento_ou_status(tmp_path):
    """Forma de Pagamento, Status de Aprovação e Status de Entrega continuam
    [H] em DAD_001 — texto livre, sem domínio fechado inventado."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Compras"]
    for coluna in (COL_COM_FORMA_PAGAMENTO, COL_COM_STATUS_APROVACAO, COL_COM_STATUS_ENTREGA):
        letra = ws.cell(row=1, column=coluna).column_letter
        dvs = [dv for dv in ws.data_validations.dataValidation if str(dv.sqref).startswith(f"{letra}2")]
        assert dvs == []


def test_valor_tem_validacao_numerica_maior_ou_igual_a_zero(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Compras"]
    letra = ws.cell(row=1, column=COL_COM_VALOR).column_letter
    dv_valor = next(dv for dv in ws.data_validations.dataValidation if str(dv.sqref).startswith(f"{letra}2"))
    assert dv_valor.type == "decimal" and dv_valor.operator == "greaterThanOrEqual"
    assert dv_valor.formula1 == "0"


def test_formula_de_planejado_busca_valor_previsto_do_servico_vinculado(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Compras"]
    formula = ws.cell(row=2, column=COL_COM_PLANEJADO).value
    assert "INDEX(Serviços!" in formula
    assert "MATCH(" in formula


def test_formula_de_variacao_reais_e_compra_menos_planejado(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Compras"]
    formula = ws.cell(row=2, column=COL_COM_VARIACAO_REAIS).value
    letra_valor = ws.cell(row=1, column=COL_COM_VALOR).column_letter
    letra_planejado = ws.cell(row=1, column=COL_COM_PLANEJADO).column_letter
    assert f"${letra_valor}2-${letra_planejado}2" in formula
    assert "Desvio" not in formula  # termo proibido (Variação, nunca Desvio)


def test_formula_de_variacao_percentual_trata_planejado_zero_como_nd(tmp_path):
    """Regra 17 — mesmo tratamento de REG-031: nunca 0%."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Compras"]
    formula = ws.cell(row=2, column=COL_COM_VARIACAO_PERCENTUAL).value
    assert '"N/D"' in formula


def test_compras_nao_aparece_em_nenhum_total_de_orcamento_ou_financeiro(tmp_path):
    """Regressão negativa: Compras não deve alimentar nenhuma fórmula de
    total de Orçamento (Etapa 3), Financeiro (Etapa 4) ou Resumo
    Financeiro — mesmo padrão de
    `test_totais_nao_somam_aportes_alteracoes_ou_custo_realizado`."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    for aba in ("Início", "Etapas", "Subetapas", "Serviços", "Financeiro", "Resumo Financeiro"):
        ws = wb[aba]
        for linha in ws.iter_rows():
            for celula in linha:
                if isinstance(celula.value, str) and celula.value.startswith("="):
                    assert "Compras" not in celula.value
                    assert "TabelaCompras" not in celula.value


# ----------------------------------------------------------------------
# D-COM-1 (Etapa 8.2) — vínculos da Compra escritos no Excel
# ----------------------------------------------------------------------
def _base_compras_com_vinculos():
    """Obra com Serviço orçado em 2.400 e três Compras: uma vinculada a
    Etapa+Subetapa+Serviço, uma sem vínculo nenhum, e duas do "mesmo
    pedido" (múltiplos itens = múltiplas linhas, REG-010/Etapa 6)."""
    from datetime import date

    from src.modelo.entidades import Compra, Etapa, Fornecedor, Obra, ServicoOrcamento, Subetapa

    base = BaseDados()
    g = base.gerador_id
    obra = base.adicionar_obra(Obra(id=g.gerar("OBRAS"), nome="Obra"))
    etapa = base.adicionar_etapa(Etapa(id=g.gerar("ETAPAS"), id_obra=obra.id, nome="Fundação"))
    subetapa = base.adicionar_subetapa(
        Subetapa(id=g.gerar("SUBETAPAS"), id_etapa=etapa.id, nome="Escavação")
    )
    servico = base.adicionar_servico(
        ServicoOrcamento(
            id=g.gerar("SERVICOS_ORCAMENTO"), id_subetapa=subetapa.id, descricao="Chapisco",
            quantidade_orcada=200.0, valor_unitario=12.0,  # Valor Previsto = 2.400
        )
    )
    forn = base.adicionar_fornecedor(
        Fornecedor(id=g.gerar("FORNECEDORES"), nome_razao_social="Fornecedor ABC")
    )
    vinculada = base.adicionar_compra(Compra(
        id=g.gerar("COMPRAS"), id_obra=obra.id, descricao="Cimento CP-II",
        id_fornecedor=forn.id, data_compra=date(2026, 3, 8), valor=2_800.0,
        id_etapa=etapa.id, id_subetapa=subetapa.id, id_servico=servico.id))
    avulsa = base.adicionar_compra(Compra(
        id=g.gerar("COMPRAS"), id_obra=obra.id, descricao="Compra avulsa",
        data_compra=date(2026, 3, 9), valor=500.0))
    item2 = base.adicionar_compra(Compra(
        id=g.gerar("COMPRAS"), id_obra=obra.id, descricao="Areia (mesmo pedido)",
        data_compra=date(2026, 3, 8), valor=300.0, id_servico=servico.id))
    return base, etapa, subetapa, servico, (vinculada, avulsa, item2)


def test_compra_vinculada_escreve_os_nomes_amigaveis_no_excel(tmp_path):
    """D-COM-1: antes da Etapa 8.2 as colunas Etapa/Subetapa/Serviço
    ficavam VAZIAS para registros existentes, e com isso o Planejado e as
    Variações (REG-010) não eram calculados no Excel."""
    base, etapa, subetapa, servico, _ = _base_compras_com_vinculos()
    caminho = tmp_path / "v.xlsx"
    construir_workbook(base).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Compras"]

    assert ws.cell(row=2, column=COL_COM_ETAPA).value == etapa.nome
    assert ws.cell(row=2, column=COL_COM_SUBETAPA).value == subetapa.nome
    assert ws.cell(row=2, column=COL_COM_SERVICO).value == servico.descricao
    # Nome amigável, nunca ID técnico (AGENTS.md §11).
    assert ws.cell(row=2, column=COL_COM_SERVICO).value != servico.id


def test_cadeia_planejado_e_variacoes_fica_calculavel_no_excel(tmp_path):
    """Com o vínculo escrito, a coluna técnica ID_Servico resolve por
    MATCH e a cadeia Planejado → Variação R$ → Variação % passa a ter
    entrada. As três continuam sendo FÓRMULA (nunca valor literal)."""
    base, _etapa, _subetapa, servico, _ = _base_compras_com_vinculos()
    caminho = tmp_path / "v.xlsx"
    construir_workbook(base).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Compras"]

    id_servico = ws.cell(row=2, column=COL_COM_ID_SERVICO).value
    assert str(id_servico).startswith("=") and "MATCH(" in id_servico
    planejado = ws.cell(row=2, column=COL_COM_PLANEJADO).value
    assert str(planejado).startswith("=") and "Serviços!" in planejado
    assert str(ws.cell(row=2, column=COL_COM_VARIACAO_REAIS).value).startswith("=")
    assert str(ws.cell(row=2, column=COL_COM_VARIACAO_PERCENTUAL).value).startswith("=")


def test_compra_sem_vinculo_continua_valida_e_com_colunas_vazias(tmp_path):
    """Regra 3 da Etapa 6 preservada: Compra pode existir sem vínculo —
    as colunas de vínculo ficam vazias, sem erro."""
    base, _etapa, _subetapa, _servico, (_v, avulsa, _i) = _base_compras_com_vinculos()
    caminho = tmp_path / "v.xlsx"
    construir_workbook(base).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Compras"]

    assert ws.cell(row=3, column=COL_COM_DESCRICAO).value == avulsa.descricao
    assert ws.cell(row=3, column=COL_COM_ETAPA).value is None
    assert ws.cell(row=3, column=COL_COM_SUBETAPA).value is None
    assert ws.cell(row=3, column=COL_COM_SERVICO).value is None


def test_multiplos_itens_do_mesmo_pedido_continuam_como_linhas_independentes(tmp_path):
    """Etapa 6: múltiplos itens = múltiplas linhas. Cada uma escreve seu
    próprio vínculo."""
    base, _etapa, _subetapa, servico, _ = _base_compras_com_vinculos()
    caminho = tmp_path / "v.xlsx"
    construir_workbook(base).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Compras"]

    assert ws.cell(row=2, column=COL_COM_SERVICO).value == servico.descricao
    assert ws.cell(row=4, column=COL_COM_SERVICO).value == servico.descricao
    assert ws.cell(row=2, column=COL_COM_DESCRICAO).value != ws.cell(row=4, column=COL_COM_DESCRICAO).value


def test_vinculos_da_compra_nao_criam_efeito_em_financeiro_orcamento_ou_execucao(tmp_path):
    """Regras 2/7/8/9 da Etapa 6 preservadas: escrever o vínculo é
    apresentação — não gera lançamento nem altera orçamento/execução."""
    from src.execucao.calculos import percentual_execucao_obra
    from src.financeiro.calculos import custo_realizado, saldo_caixa
    from src.orcamento.calculos import total_previsto_obra

    base, _etapa, _subetapa, _servico, _ = _base_compras_com_vinculos()
    obra_id = next(iter(base.obras))
    assert base.financeiro == {}
    assert custo_realizado(base, obra_id) == 0
    assert saldo_caixa(base, obra_id) == 0
    assert total_previsto_obra(base, obra_id) == 2_400.0  # só o orçamento do serviço
    assert percentual_execucao_obra(base, obra_id) == 0.0
