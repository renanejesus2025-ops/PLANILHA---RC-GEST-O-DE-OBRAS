"""
Testes das abas Financeiro / Pagamentos / Resumo Financeiro (Etapa 4).

Como não há motor de planilha neste ambiente (ver relatórios das Etapas
2/3), a verificação é sobre o TEXTO das fórmulas — não sobre um valor
calculado — exatamente como `testes/test_excel_orcamento.py`.
"""

from __future__ import annotations

from datetime import date

import openpyxl
import pytest

from src.base_dados.repositorio import BaseDados
from src.excel.construtor_workbook import (
    COL_FIN_A_PAGAR,
    COL_FIN_APORTE,
    COL_FIN_CUSTO_REALIZADO,
    COL_FIN_ENTRADA_SAIDA,
    COL_FIN_ID_OBRA,
    COL_FIN_TOTAL_PAGO,
    COL_PAG_ID_FINANCEIRO,
    COLUNAS_FINANCEIRO,
    COLUNAS_PAGAMENTOS,
    construir_workbook,
)
from src.modelo.entidades import Financeiro, Obra, Pagamento
from src.modelo.enums import TipoLancamentoFinanceiro


def _base_com_obra(base: BaseDados | None = None) -> tuple[BaseDados, Obra]:
    base = base or BaseDados()
    obra = base.adicionar_obra(Obra(id=base.gerador_id.gerar("OBRAS"), nome="Obra"))
    return base, obra


# --------------------------------------------------------------------
# Estrutura das novas abas
# --------------------------------------------------------------------
def test_abas_financeiro_pagamentos_resumo_existem(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    assert {"Financeiro", "Pagamentos", "Resumo Financeiro"} <= set(wb.sheetnames)


@pytest.mark.parametrize(
    "aba,colunas",
    [("Financeiro", COLUNAS_FINANCEIRO), ("Pagamentos", COLUNAS_PAGAMENTOS)],
)
def test_cabecalhos_financeiro_pagamentos(tmp_path, aba, colunas):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)[aba]
    titulos_esperados = [titulo for titulo, _largura, _oculta, _calculada in colunas]
    titulos_no_arquivo = [ws.cell(row=1, column=i + 1).value for i in range(len(colunas))]
    assert titulos_no_arquivo == titulos_esperados


def test_colunas_tecnicas_de_id_ficam_ocultas_financeiro_pagamentos(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    for aba, colunas in [("Financeiro", COLUNAS_FINANCEIRO), ("Pagamentos", COLUNAS_PAGAMENTOS)]:
        ws = wb[aba]
        for indice, (_titulo, _largura, oculta, _calculada) in enumerate(colunas, start=1):
            letra = ws.cell(row=1, column=indice).column_letter
            assert ws.column_dimensions[letra].hidden is oculta


# --------------------------------------------------------------------
# Seção 8 — dropdown de Tipo com os 4 valores homologados
# --------------------------------------------------------------------
def test_dropdown_de_tipo_usa_exatamente_os_4_rotulos_homologados(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Financeiro"]
    dvs_tipo = [dv for dv in ws.data_validations.dataValidation if str(dv.sqref).startswith("B")]
    assert len(dvs_tipo) == 1
    formula = dvs_tipo[0].formula1
    for tipo in TipoLancamentoFinanceiro:
        assert tipo.rotulo in formula
    assert formula.count(",") == 3  # exatamente 4 itens


# --------------------------------------------------------------------
# Seção 9/10/11/12 — colunas derivadas por Tipo
# --------------------------------------------------------------------
def test_coluna_aporte_isola_apenas_lancamentos_do_tipo_aporte(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Financeiro"]
    formula = ws.cell(row=2, column=COL_FIN_APORTE).value
    assert formula == '=IF($B2="Aporte",$E2,"")'


def test_coluna_custo_realizado_isola_apenas_despesa_custo(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Financeiro"]
    formula = ws.cell(row=2, column=COL_FIN_CUSTO_REALIZADO).value
    assert formula == '=IF($B2="Despesa/Custo",$E2,"")'


def test_coluna_a_pagar_e_custo_menos_total_pago(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Financeiro"]
    formula = ws.cell(row=2, column=COL_FIN_A_PAGAR).value
    assert formula == '=IF($B2="Despesa/Custo",$H2-$I2,"")'


def test_coluna_total_pago_soma_pagamentos_pela_aba_pagamentos(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Financeiro"]
    formula = ws.cell(row=2, column=COL_FIN_TOTAL_PAGO).value
    assert "SUMIFS(" in formula
    assert "Pagamentos!$D:$D" in formula  # Valor Pago
    assert "Pagamentos!$F:$F" in formula  # ID_Financeiro (vínculo técnico)


def test_coluna_entrada_saida_nunca_classifica_despesa_como_entrada_ou_saida_direta(tmp_path):
    """Seção 16: o reconhecimento do custo não é, por si só, Entrada nem Saída de Caixa."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Financeiro"]
    formula = ws.cell(row=2, column=COL_FIN_ENTRADA_SAIDA).value
    assert '"Despesa/Custo","Entrada"' not in formula
    assert '"Despesa/Custo","Saída"' not in formula


# --------------------------------------------------------------------
# Seção 20 — Pagamentos vinculados por Descrição (nome amigável), nunca por ID
# --------------------------------------------------------------------
def test_dropdown_de_lancamento_em_pagamentos_aponta_para_descricao_nao_para_id(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    lista = wb.defined_names["Lista_Financeiro_Descricoes"].attr_text
    assert "Financeiro!$D$" in lista
    assert "$A$" not in lista


def test_id_financeiro_em_pagamentos_e_sempre_formula(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Pagamentos"]
    formula = ws.cell(row=2, column=COL_PAG_ID_FINANCEIRO).value
    assert "INDEX(Financeiro!$A:$A" in formula
    assert "MATCH($B2,Financeiro!$D:$D,0)" in formula


def test_pagamento_literal_mostra_descricao_amigavel_do_lancamento(tmp_path):
    base, obra = _base_com_obra()
    despesa = base.adicionar_financeiro(
        Financeiro(
            id=base.gerador_id.gerar("FINANCEIRO"),
            id_obra=obra.id,
            tipo=TipoLancamentoFinanceiro.DESPESA,
            data=date(2026, 9, 17),
            valor=10_000,
            descricao="Compra de cimento",
        )
    )
    base.adicionar_pagamento(
        Pagamento(id=base.gerador_id.gerar("PAGAMENTOS"), id_financeiro=despesa.id, data=date(2026, 9, 17), valor=4_000)
    )
    caminho = tmp_path / "v.xlsx"
    construir_workbook(base).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Pagamentos"]
    assert ws.cell(row=2, column=2).value == "Compra de cimento"  # nunca o ID técnico


# --------------------------------------------------------------------
# Seção 27 — formatação monetária brasileira "R$ 0,00"
# --------------------------------------------------------------------
def test_colunas_monetarias_de_financeiro_usam_formato_r_dollar(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Financeiro"]
    for coluna in (5, COL_FIN_APORTE, COL_FIN_CUSTO_REALIZADO, COL_FIN_TOTAL_PAGO, COL_FIN_A_PAGAR):
        assert "R$" in ws.cell(row=2, column=coluna).number_format


def test_resumo_financeiro_usa_formato_r_dollar(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Resumo Financeiro"]
    assert "R$" in ws.cell(row=2, column=2).number_format  # Orçamento Inicial


# --------------------------------------------------------------------
# Seção 21 — Valor sem sinal negativo (validação numérica >= 0)
# --------------------------------------------------------------------
def test_validacao_de_valor_em_financeiro_e_pagamentos_e_maior_ou_igual_a_zero(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)

    ws_fin = wb["Financeiro"]
    dv_fin = next(dv for dv in ws_fin.data_validations.dataValidation if str(dv.sqref).startswith("E2"))
    assert dv_fin.type == "decimal" and dv_fin.operator == "greaterThanOrEqual" and dv_fin.formula1 == "0"

    ws_pag = wb["Pagamentos"]
    dv_pag = next(dv for dv in ws_pag.data_validations.dataValidation if str(dv.sqref).startswith("D2"))
    assert dv_pag.type == "decimal" and dv_pag.operator == "greaterThanOrEqual" and dv_pag.formula1 == "0"


# --------------------------------------------------------------------
# Resumo Financeiro — fórmulas centrais (Seções 4/5/6/7/15)
# --------------------------------------------------------------------
def test_resumo_orcamento_vigente_soma_inicial_e_aportes(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Resumo Financeiro"]
    assert ws.cell(row=2, column=2).value == "=Início!$B$8"
    assert ws.cell(row=3, column=2).value.startswith("=SUM(Financeiro!$G$2:$G$")
    assert ws.cell(row=4, column=2).value == "=B2+B3"


def test_resumo_percentual_consumido_trata_divisao_por_zero(tmp_path):
    """Seção 7: nunca 0% quando o Orçamento Vigente é zero — usar 'N/D'."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Resumo Financeiro"]
    formula = ws.cell(row=7, column=2).value
    assert formula == '=IF(B4=0,"N/D",B5/B4)'


def test_resumo_saldo_de_caixa_nao_usa_custo_realizado_diretamente(tmp_path):
    """Seção 16: Saldo de Caixa usa Pagamentos, nunca a Despesa/Custo em si."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Resumo Financeiro"]
    for linha in range(1, 20):
        rotulo = ws.cell(row=linha, column=1).value
        formula = ws.cell(row=linha, column=2).value
        if rotulo in ("Total de Entradas", "Total de Saídas", "Saldo de Caixa") and isinstance(formula, str):
            assert "Financeiro!$H" not in formula  # coluna Custo Realizado nunca referenciada aqui


def test_resumo_estouro_de_orcamento_nao_bloqueia_apenas_informa(tmp_path):
    """Seção 29: apenas indicador textual SIM/NÃO — sem formatação condicional/bloqueio."""
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Resumo Financeiro"]
    linhas_rotulo = {ws.cell(row=r, column=1).value: r for r in range(1, 20)}
    linha_estouro = linhas_rotulo["Estouro de Orçamento (indicador Operador-only)"]
    formula = ws.cell(row=linha_estouro, column=2).value
    assert formula == '=IF(B6<0,"SIM","NÃO")'
    assert len(ws.conditional_formatting._cf_rules) == 0


# --------------------------------------------------------------------
# ID_Obra técnico de Financeiro (vínculo oculto, mesmo padrão das demais abas)
# --------------------------------------------------------------------
def test_id_obra_de_financeiro_referencia_base_dados(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    ws = openpyxl.load_workbook(caminho)["Financeiro"]
    formula = ws.cell(row=2, column=COL_FIN_ID_OBRA).value
    assert "Base_Dados!$B$2" in formula


# --------------------------------------------------------------------
# Regressão explícita (Seção 40/41): nenhum termo "Desvio" foi introduzido
# --------------------------------------------------------------------
def test_nenhum_termo_desvio_foi_introduzido_nas_abas_novas(tmp_path):
    caminho = tmp_path / "v.xlsx"
    construir_workbook(BaseDados()).save(caminho)
    wb = openpyxl.load_workbook(caminho)
    for aba in ("Financeiro", "Pagamentos", "Resumo Financeiro"):
        ws = wb[aba]
        for linha in ws.iter_rows():
            for celula in linha:
                if isinstance(celula.value, str):
                    assert "Desvio" not in celula.value and "desvio" not in celula.value
